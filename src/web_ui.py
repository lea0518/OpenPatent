import sys
sys.stdout.reconfigure(encoding='utf-8')
import gradio as gr
from docx import Document
import numpy as np
import time
from pdf_processor import PDFProcessor
from vector_db import VectorDB
from llm_integration import PatentGenerator
import os
import evaluator
# os.environ["SSL_CERT_FILE"] = r"H:\anadonda\envs\OpenPatent\Library\ssl\cacert.pem"
from docx.oxml.ns import qn

# 【阶段4】闭环控制参数
ISSUE_THRESHOLD = 2        # 问题总数 <= 此值即达标停止（阶段2基线约2，达到基线水平即收手，不为清零而做高风险重写）
MAX_REWRITE_ROUNDS = 2     # 最多自动重写轮数，防止无限循环、控制 API 开销

class WebUI:
    """
    网页用户界面类，用于创建和管理专利生成系统的用户界面。
    """
    def __init__(self):
        """
        初始化网页用户界面。
        """
        self.patent_generator = PatentGenerator()
        self.current_doc_type = None
        self.tech_doc_name = None          # 记住上传的技术文档名（不含扩展名），用于命名输出
        self.db_paths = {
            # 用 os.path.join 而非反斜杠字面量，保证 Linux/macOS 下也能正确解析
            '摘要': os.path.join('dbs', 'abstract'),
            '说 明 书': os.path.join('dbs', 'specification'),
            '权 利 要 求 书': os.path.join('dbs', 'claims'),
        }
        self.use_existing_db = False

    def init_interface(self):
        """
        初始化用户界面。

        返回:
        gr.Blocks: Gradio 界面块。
        """
        with gr.Blocks(title="OpenPatent 专利生成系统") as demo:
            gr.Markdown("## OpenPatent 专利文档生成系统")
            
            with gr.Tab("1. 选择参考专利"):
                ref_patents = gr.Files(label="上传参考专利文件(PDF)")
                process_btn = gr.Button("处理专利文件")
                process_btn2 = gr.Button("已有本地知识库，点击这里")
                process_output = gr.Markdown()
            
            with gr.Tab("2. 上传技术文档"):
                tech_doc = gr.File(label="技术文档(.docx)")
                
            with gr.Tab("3. 生成专利文档"):
                stage_selector = gr.Dropdown(
                    choices=[
                        ("阶段0 基线（无术语表, 高温, 独立生成）", 0),
                        ("阶段1 术语约束（术语表 + 低温）", 1),
                        ("阶段2 分层串联（+ 权利要求/摘要依赖说明书）", 2),
                        ("阶段4 生成-评估-重写闭环（+ 自动评估并重写至达标）", 3),
                    ],
                    value=2,
                    label="消融实验阶段（切换后点一键分层生成即按该阶段运行）",
                )
                with gr.Row():
                    gen_all_btn = gr.Button("一键分层生成（说明书→权利要求→摘要）", variant="primary")
                with gr.Row():
                    gen_spec_btn = gr.Button("生成说明书")
                    gen_abstract_btn = gr.Button("生成摘要")
                    gen_claims_btn = gr.Button("生成权利要求书")
                
                with gr.Column():
                    output_preview = gr.Chatbot(label="专利生成过程", height=500, elem_id="centered-chat")
                
                with gr.Row():
                    user_feedback = gr.Textbox(label="修改意见", lines=3)
                    submit_feedback = gr.Button("提交反馈", variant="primary")
            
            # 绑定事件
            process_btn.click(self.process_patents, inputs=ref_patents, outputs=process_output)
            process_btn2.click(self.load_existing_db, outputs=process_output)
            gen_all_btn.click(self.generate_all_layered, inputs=[tech_doc, stage_selector], outputs=output_preview)
            # 三个单独按钮同样传入 stage_selector：否则它们会沿用上一次一键生成残留的
            # stage 值，导致 UI 显示的档位与实际运行档位不一致，静默产出错误实验数据
            gen_spec_btn.click(self.generate_specification, inputs=[tech_doc, stage_selector], outputs=output_preview)
            gen_abstract_btn.click(self.generate_abstract, inputs=[tech_doc, stage_selector], outputs=output_preview)
            gen_claims_btn.click(self.generate_claims, inputs=[tech_doc, stage_selector], outputs=output_preview)
            submit_feedback.click(self.submit_feedback, inputs=user_feedback, outputs=output_preview)
            
        return demo

    def process_patents(self, files):
        """
        处理上传的参考专利文件，创建向量数据库并保存索引。

        参数:
        files (list): 上传的参考专利文件列表。

        返回:
        str: 处理结果信息。
        """
        if not files:
            return "请上传参考专利文件"
        processor = PDFProcessor()

        section_array = []
        skipped = []
        missing = []          # 记录"解析成功但缺某章节"的情况，便于用户排查
        for file in files:
            try:
                sections = processor.split_pdf(file.name)
            except Exception as e:
                skipped.append(f"{os.path.basename(file.name)}（{type(e).__name__}）")
                continue
            for db_type in self.db_paths:
                section_content = sections.get(db_type)
                # 缺章节时统一存空串而非 None：None 送进 embedding API 会被 400 拒绝，
                # 导致整批建库失败。空串由 create_index 负责跳过。
                if not section_content or not section_content.strip():
                    missing.append(f"{os.path.basename(file.name)} 缺「{db_type}」")
                    section_array.append("")
                else:
                    section_array.append(section_content)
        if not section_array:
            return "所有上传文件都无法解析，请检查 PDF 是否为标准专利文本格式（非扫描件）"
        section_array = np.array(section_array, dtype=object)
        section_array = section_array.reshape((-1,3))
        section_array = section_array.T
        self.db_list = [0 for _ in range(3)]
        try:
            for i, db_type in enumerate(self.db_paths):
                self.db_list[i] = VectorDB(db_type)
                self.db_list[i].create_index(section_array[i])
                self.db_list[i].save_index(self.db_paths[db_type])
        except Exception as e:
            self.use_existing_db = False
            return f"❌ 建立知识库失败：{type(e).__name__}: {e}\n\n知识库未就绪，请修正后重试。"
        self.use_existing_db = True
        msg = "参考专利处理完成，已建立三个知识库！"
        if skipped:
            msg += "\n\n⚠️ 以下文件损坏或无法解析，已跳过：\n" + "\n".join(skipped)
        if missing:
            msg += "\n\n⚠️ 以下章节未能切出（该条已跳过，不影响其余）：\n" + "\n".join(missing)
        return msg

    def load_existing_db(self):
        """
        加载已有的本地知识库。

        返回:
        str: 加载结果信息。
        """
        db_list = [0 for _ in range(3)]
        try:
            for i, db_type in enumerate(self.db_paths):
                db_list[i] = VectorDB(db_type)
                db_list[i].load_index(self.db_paths[db_type])
        except Exception as e:
            # 加载失败不能把 use_existing_db 置真，否则会带着空库去生成（RAG 静默空转）
            self.use_existing_db = False
            return (f"❌ 加载本地知识库失败：{type(e).__name__}: {e}\n\n"
                    f"请点击上方『处理专利文件』重新建库。")
        # 全部成功后才提交状态，避免部分加载成功导致 db_list 半残
        self.db_list = db_list
        self.use_existing_db = True
        total = sum(db.index.ntotal for db in self.db_list)
        return f"已加载本地知识库（三库合计 {total} 条向量）"

    def generate_specification(self, tech_doc, stage=2):
        """
        生成专利说明书。

        参数:
        tech_doc: 上传的技术文档。
        stage (int): 消融档位，由 UI 下拉框传入。

        返回:
        list: 包含系统消息和生成内容的列表。
        """
        return self._generate_draft(tech_doc, "说 明 书", "说明书", stage)

    def generate_abstract(self, tech_doc, stage=2):
        """
        生成专利摘要。

        参数:
        tech_doc: 上传的技术文档。
        stage (int): 消融档位，由 UI 下拉框传入。

        返回:
        list: 包含系统消息和生成内容的列表。
        """
        return self._generate_draft(tech_doc, "摘要", "摘要", stage)

    def generate_claims(self, tech_doc, stage=2):
        """
        生成专利权利要求书。

        参数:
        tech_doc: 上传的技术文档。
        stage (int): 消融档位，由 UI 下拉框传入。

        返回:
        list: 包含系统消息和生成内容的列表。
        """
        return self._generate_draft(tech_doc, "权 利 要 求 书", "权利要求书", stage)

    def generate_all_layered(self, tech_doc, stage=2):
        """
        【消融实验】一键分层生成：按 说明书 → 权利要求 → 摘要 的顺序串联生成。
        stage 控制启用哪些改进（0基线 / 1术语约束 / 2分层串联），供消融对比。
        阶段>=2 时，后层以已生成的说明书为事实依据，保证特征支撑与跨部分一致。
        """
        if not self.use_existing_db:
            return [("系统", "请先处理参考专利或选择已有知识库")]
        if tech_doc is None:
            return [("系统", "请先上传技术文档")]

        stage = int(stage)
        # 把档位下发给生成器，generate_draft 会据此决定温度/术语表/串联依赖
        self.patent_generator.stage = stage

        # 读取技术文档
        doc = Document(tech_doc.name)
        # 记住上传文档名（去掉路径和扩展名），供保存时命名用
        self.tech_doc_name = os.path.splitext(os.path.basename(tech_doc.name))[0]
        query = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])

        # 每次一键生成前重置状态，避免上一份文档的草稿/术语表/检索上下文污染（多文档实验必需）
        self.patent_generator.current_draft = {}
        self.patent_generator.glossary = ""
        self.patent_generator.contexts = {}

        # 术语表仅在 stage>=1 构建（stage0 基线不用术语表）
        if stage >= 1:
            glossary = self.patent_generator.build_glossary(query)
            print(f"【术语表】\n{glossary}")
            try:
                with open("glossary_latest.txt", "w", encoding="utf-8") as gf:
                    gf.write(glossary)
            except Exception as e:
                print(f"术语表存档失败: {e}")

        db_map = {"摘要": 0, "说 明 书": 1, "权 利 要 求 书": 2}
        messages = [("系统", f"【阶段{stage}】开始生成：说明书 → 权利要求 → 摘要 ...")]

        # 顺序固定：说明书先生成，stage>=2 时权利要求和摘要会依赖它
        plan = [
            ("说 明 书", "说明书"),
            ("权 利 要 求 书", "权利要求书"),
            ("摘要", "摘要"),
        ]
        for db_type, doc_type in plan:
            vector_db = self.db_list[db_map[db_type]]
            related = vector_db.query(query, top_k=2)
            context = "\n".join(related) if related else "无相关专利内容"
            content = self.patent_generator.generate_draft(query, context, doc_type)
            self.current_doc_type = doc_type
            messages.append(("助手", f"【{doc_type}】\n{content}"))

        # 【阶段4】stage>=3：接入 生成→评估→重写 自动闭环（query 此处仍是交底书全文）
        if stage >= 3:
            messages = self._run_closed_loop(query, messages)

        return messages

    # ===================== 【阶段4】生成→评估→重写 自动闭环 =====================
    def _strip_tags(self, content: str) -> str:
        """把带 <标题>/<段落> 标签的生成内容剥成纯文本，喂给评估器。
        与 _save_content_to_docx 同款正则，保证与阶段3离线评估同口径。剥出为空则回退原文。"""
        import re
        parts = [t.strip() for _, t in
                 re.findall(r'<(标题|段落)>(.*?)</\1>', content or "", flags=re.S) if t.strip()]
        text = "\n".join(parts)
        return text if text else (content or "")

    def _build_feedback_by_part(self, issues: dict) -> dict:
        """把评估器返回的结构化问题，按"重写目标路由"拆成各部分的针对性反馈。
        返回 {"说明书":文本或None, "权利要求书":..., "摘要":...}，无问题的部分为 None。
        路由：无支撑特征→仅权利要求；术语漂移→权利要求+摘要；自造术语→三部分。"""
        drift = issues.get("term_drift", {}).get("items", []) or []
        unsup = issues.get("unsupported_claims", {}).get("items", []) or []
        fab = issues.get("fabricated_terms", {}).get("items", []) or []

        def blk(title, tip, items):
            if not items:
                return ""
            return "\n".join([f"【{title}】{tip}"] + [f"- {it}" for it in items]) + "\n"

        head = ("以下是专利审查发现的问题，请【只针对这些问题】修改本部分，"
                "不要改动无关内容，保持原有 <标题>/<段落> XML 标签格式。\n")
        fb = {"说明书": None, "权利要求书": None, "摘要": None}

        claim_body = (blk("无说明书支撑的权利要求特征", "（必须删除，或改写为说明书中已有的特征）：", unsup)
                      + blk("术语漂移", "（必须统一为说明书中的写法）：", drift)
                      + blk("自造术语", "（必须删除或替换为交底书/术语表中的规范术语）：", fab))
        if claim_body:
            fb["权利要求书"] = head + claim_body

        abst_body = (blk("术语漂移", "（必须统一为说明书中的写法）：", drift)
                     + blk("自造术语", "（必须删除或替换为规范术语）：", fab))
        if abst_body:
            fb["摘要"] = head + abst_body

        spec_body = blk("自造术语", "（说明书为事实源，必须删除或替换为规范术语）：", fab)
        if spec_body:
            fb["说明书"] = head + spec_body

        return fb

    def _run_closed_loop(self, tech, messages):
        """闭环主体：初评 → 若问题>阈值则按部分重写 → 复评，最多 MAX_REWRITE_ROUNDS 轮。
        tech 为交底书全文。全程把 total_issues 轨迹记入 messages 与 closedloop_latest.json。"""
        gen = self.patent_generator
        glossary = gen.glossary

        def evaluate_now():
            spec = self._strip_tags(gen.current_draft.get("说明书", ""))
            abst = self._strip_tags(gen.current_draft.get("摘要", ""))
            claim = self._strip_tags(gen.current_draft.get("权利要求书", ""))
            return evaluator.extract_issues(tech, spec, abst, claim, glossary)

        def total_of(iss):
            return iss.get("total_issues") if "_parse_error" not in iss else None

        def summarize(iss, label):
            if "_parse_error" in iss:
                return f"{label}：评估解析失败（{iss.get('_parse_error')}），闭环中止。"
            return (f"{label}：问题总数 {iss['total_issues']}"
                    f"（术语漂移 {iss['term_drift']['count']} / 无支撑权利要求 "
                    f"{iss['unsupported_claims']['count']} / 自造术语 {iss['fabricated_terms']['count']}）")

        trajectory, rounds_log = [], []
        issues = evaluate_now()
        t = total_of(issues)
        messages.append(("系统", "【阶段4闭环】" + summarize(issues, "初评")))
        if t is None:
            self._dump_closedloop(trajectory, rounds_log)
            return messages
        trajectory.append(t)
        rounds_log.append({"round": 0, "issues": issues})

        rnd = 0
        while t > ISSUE_THRESHOLD and rnd < MAX_REWRITE_ROUNDS:
            rnd += 1
            # 回滚快照：重写前存一份当前草稿。若本轮改完问题反而变多，就还原，避免"越改越糟"
            snapshot = dict(gen.current_draft)
            prev_t = t
            fbmap = self._build_feedback_by_part(issues)
            # 固定顺序：先修事实源(说明书)，下游再对齐更新后的说明书
            for doc_type in ("说明书", "权利要求书", "摘要"):
                if fbmap.get(doc_type):
                    revised = gen.revise_draft(fbmap[doc_type], doc_type)
                    self.current_doc_type = doc_type
                    messages.append(("助手", f"【第{rnd}轮重写·{doc_type}】\n{revised}"))
            issues = evaluate_now()
            t2 = total_of(issues)
            messages.append(("系统", summarize(issues, f"第{rnd}轮重写后")))
            if t2 is None:
                break
            trajectory.append(t2)
            rounds_log.append({"round": rnd, "issues": issues})
            # 本轮反而变差 → 回滚到重写前，并停止（既然改不动，再改只会更糟）
            if t2 > prev_t:
                gen.current_draft = snapshot
                messages.append(("系统", f"⚠️ 第{rnd}轮重写后问题数从 {prev_t} 升至 {t2}，"
                                         f"已回滚到本轮重写前的版本并停止闭环。"))
                trajectory.append(prev_t)  # 记录回滚后的最终问题数
                break
            t = t2

        messages.append(("系统", f"【阶段4闭环结束】问题数轨迹: {trajectory}"
                                 f"（阈值 {ISSUE_THRESHOLD}，最多 {MAX_REWRITE_ROUNDS} 轮）"))
        self._dump_closedloop(trajectory, rounds_log)
        # 自动归档本轮全部产物：闭环跑一次成本很高，若只写固定名的 latest 文件，
        # 下次运行即被覆盖，事后无法用 term_metrics.py 复算指标、也无法复核文本改动
        archive = self._archive_run(trajectory, rounds_log, glossary)
        if archive:
            messages.append(("系统", f"📦 本轮产物已归档到：{archive}"))
        return messages

    def _archive_run(self, trajectory, rounds_log, glossary):
        """把本次闭环运行的全部产物落盘到带时间戳的目录，供论文复算与复核。

        归档内容与 experiments/stageX_*/outputs/ 目录结构对齐：
          三份 docx + closedloop.json + glossary.txt
        返回归档目录路径；失败返回 None（不影响主流程）。
        """
        import json
        try:
            base = self.tech_doc_name or "patent"
            out_dir = os.path.join("experiments", "stage4_closedloop",
                                   f"{base}_{int(time.time())}")
            os.makedirs(out_dir, exist_ok=True)
            for doc_type, content in self.patent_generator.current_draft.items():
                self._save_content_to_docx(content, doc_type, out_dir)
            with open(os.path.join(out_dir, "closedloop.json"), "w", encoding="utf-8") as f:
                json.dump({"stage": 3, "threshold": ISSUE_THRESHOLD,
                           "max_rounds": MAX_REWRITE_ROUNDS,
                           "trajectory": trajectory, "rounds": rounds_log},
                          f, ensure_ascii=False, indent=2)
            with open(os.path.join(out_dir, "glossary.txt"), "w", encoding="utf-8") as f:
                f.write(glossary or "")
            return out_dir
        except Exception as e:
            print(f"闭环产物归档失败: {e}")
            return None

    def _dump_closedloop(self, trajectory, rounds_log):
        """把闭环问题数轨迹落盘，供论文取数。"""
        import json
        try:
            with open("closedloop_latest.json", "w", encoding="utf-8") as f:
                json.dump({"stage": 3, "threshold": ISSUE_THRESHOLD,
                           "max_rounds": MAX_REWRITE_ROUNDS,
                           "trajectory": trajectory, "rounds": rounds_log},
                          f, ensure_ascii=False, indent=2)
        except Exception as e:
            print(f"闭环日志写入失败: {e}")

    def _generate_draft(self, tech_doc, db_type: str, doc_type: str, stage=2):
        """
        生成专利文档初稿。

        参数:
        tech_doc: 上传的技术文档。
        db_type (str): 数据库类型，如 "摘要", "说明书", "权利要求书"。
        doc_type (str): 文档类型，如 "说明书", "摘要", "权利要求书"。
        stage (int): 消融档位，由 UI 下拉框传入，避免沿用上次残留值。

        返回:
        list: 包含系统消息和生成内容的列表。
        """
        if not self.use_existing_db:
            return [("系统", "请先处理参考专利或选择已有知识库")]
        if tech_doc is None:
            return [("系统", "请先上传技术文档")]

        # 与一键生成保持一致：以 UI 当前档位为准，否则会沿用上次残留的 stage
        stage = int(stage)
        self.patent_generator.stage = stage

        # 读取技术文档内容
        doc = Document(tech_doc.name)
        # 记住上传文档名（去掉路径和扩展名），供保存时命名用
        new_name = os.path.splitext(os.path.basename(tech_doc.name))[0]
        # 换了交底书就必须清空上一份的草稿与术语表，否则会用 A 文档的术语表去写 B 文档
        if self.tech_doc_name != new_name:
            self.patent_generator.current_draft = {}
            self.patent_generator.glossary = ""
            self.patent_generator.contexts = {}
        self.tech_doc_name = new_name
        query = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])

        # 【阶段1】首次生成时，用技术文档构建术语表，供后续各部分共享，保证术语一致
        # （stage0 基线不用术语表）
        if stage >= 1 and not self.patent_generator.glossary:
            glossary = self.patent_generator.build_glossary(query)
            print(f"【术语表】\n{glossary}")
            # 自动存档术语表，供后续指标计算使用（不依赖翻终端）
            try:
                with open("glossary_latest.txt", "w", encoding="utf-8") as gf:
                    gf.write(glossary)
                print("【术语表已存到 glossary_latest.txt】")
            except Exception as e:
                print(f"术语表存档失败: {e}")

        # 加载向量数据库
        db_map = {"摘要":0,"说 明 书":1,"权 利 要 求 书":2}
        vector_db = self.db_list[db_map[db_type]]
        # 检索相关内容
        related_patents = vector_db.query(query, top_k=2)
        print(related_patents)
        context = "\n".join(related_patents) if related_patents else "无相关专利内容"
        
        # 生成专利文档
        try:
            content = self.patent_generator.generate_draft(query, context, doc_type)
            self.current_doc_type = doc_type
            return [
                ("系统", "开始生成专利文档..."),
                ("助手", content)
            ]
        except Exception as e:
            return [
                ("系统", "开始生成专利文档..."),
                ("助手", f"生成失败: {str(e)}")
            ]

    def _save_content_to_docx(self, content, doc_type, out_dir="."):
        """把一份草稿内容按专利格式存成 docx，存到 out_dir，返回文件路径。"""
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_LINE_SPACING

        doc = Document()
        doc.styles['Normal'].font.name = '宋体'
        doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        heading_style = doc.styles['Heading 1']
        heading_style.font.name = '宋体'
        heading_style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        heading_style.font.size = Pt(10.5)
        heading_style.font.bold = True
        heading_style.paragraph_format.space_before = Pt(6)
        heading_style.paragraph_format.space_after = Pt(6)
        heading_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

        body_style = doc.styles['Normal']
        body_style.font.name = '宋体'
        body_style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        body_style.font.size = Pt(10.5)
        body_style.paragraph_format.space_before = Pt(0)
        body_style.paragraph_format.space_after = Pt(0)
        body_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        body_style.paragraph_format.first_line_indent = Pt(0)

        import re
        # 只提取成对完整的标签，忽略截断残片（如 "...</parag"）
        for tag, text in re.findall(r'<(标题|段落)>(.*?)</\1>', content, flags=re.S):
            text = text.strip()
            if not text:
                continue
            if tag == '标题':
                doc.add_paragraph(text, style='Heading 1')
            else:
                doc.add_paragraph(text)

        base = self.tech_doc_name or "patent"
        filename = os.path.join(out_dir, f"{base}_{doc_type}.docx")
        doc.save(filename)
        return filename

    def submit_feedback(self, feedback):
        """
        提交用户反馈，根据反馈内容进行文档保存或修订。

        参数:
        feedback (str): 用户的反馈意见。

        返回:
        list: 包含系统消息和处理结果的列表。
        """
        if not self.current_doc_type:
            return [("系统", "请先生成草案")]
        if not feedback.strip():
            return [("系统", "请输入修改意见")]

        try:
            if '满意' in feedback:
                messages = [("系统", "文档已确认满意，开始保存...")]
                drafts = self.patent_generator.current_draft
                if not drafts:
                    return [("系统", "没有可保存的草稿")]
                # 本批三份文件统一放进一个子文件夹：outputs/文档名_时间戳/
                base = self.tech_doc_name or "patent"
                out_dir = os.path.join("outputs", f"{base}_{int(time.time())}")
                os.makedirs(out_dir, exist_ok=True)
                # 遍历所有已生成的部分（说明书/权利要求书/摘要），逐个保存
                for doc_type, content in drafts.items():
                    filename = self._save_content_to_docx(content, doc_type, out_dir)
                    messages.append(("助手", f"【{doc_type}】已保存为：{filename}"))
                messages.append(("系统", f"全部保存在文件夹：{out_dir}"))
                return messages
            else:
                revised_content = self.patent_generator.revise_draft(feedback, self.current_doc_type)
                messages = [
                    ("系统", "开始修订文档..."),
                    ("助手", revised_content)
                ]
                return messages
        except Exception as e:
            return [
                ("系统", "处理反馈时发生错误"),
                ("助手", f"错误详情: {str(e)}")
            ]

if __name__ == "__main__":
    WebUI().init_interface().launch(show_api=False)
